#!/usr/bin/env python3
"""Generate an ND30 DOCX from a format profile and a canonical content spec.

Usage:
    python generate_from_profile_and_content.py profile.json content.json output.docx --mode auto
    python generate_from_profile_and_content.py profile.json content.json output.docx --mode replicate --source-docx source.docx
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from compose_from_source_docx import clear_body_preserve_sections
from create_nd30_docx import Spec, build, optimize_pagination, set_doc_defaults
from nd30_docx_tools import clone_patch_docx, content_to_placeholder_map
from validate_nd30_spec import validate as validate_word


DOC_TEMPLATE_MAP = {
    'cong-van': 'cong-van-template.docx',
    'quyet-dinh-truc-tiep': 'quyet-dinh-template.docx',
    'quyet-dinh-gian-tiep': 'quyet-dinh-template.docx',
    'giay-moi': 'giay-moi-template.docx',
    'bien-ban': 'bien-ban-template.docx',
}


def load_json(path: Path) -> dict:
    data = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(data, dict):
        raise ValueError('JSON root must be an object')
    return data


def find_template(document_type: str, profile: dict, template_dir: Path) -> Path | None:
    hint = str(profile.get('source_document', '') or profile.get('template_hint', '')).strip()
    if hint:
        candidate = template_dir / hint
        if candidate.exists():
            return candidate
    mapped = DOC_TEMPLATE_MAP.get(document_type)
    if mapped:
        candidate = template_dir / mapped
        if candidate.exists():
            return candidate
    return None


def generate_canonical(content: dict, output: Path) -> None:
    spec = Spec(content)
    doc = Document()
    set_doc_defaults(doc)
    build(doc, spec)
    optimize_pagination(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def generate_replicate_shell(source_docx: Path, content: dict, output: Path) -> None:
    spec = Spec(content)
    doc = Document(source_docx)
    clear_body_preserve_sections(doc)
    build(doc, spec)
    optimize_pagination(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def choose_replicate_strategy(profile: dict, content: dict, source_docx: Path, requested: str) -> str:
    if requested in {'clone-patch', 'shell-rebuild'}:
        return requested
    if profile.get('placeholder_keys'):
        return 'clone-patch'
    if isinstance(content.get('placeholder_values'), dict) and content.get('placeholder_values'):
        return 'clone-patch'
    if source_docx and 'template' in source_docx.name.lower():
        return 'clone-patch'
    return 'shell-rebuild'


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('profile_json')
    parser.add_argument('content_json')
    parser.add_argument('output_docx')
    parser.add_argument('--mode', choices=['auto', 'canonical', 'replicate'], default='auto')
    parser.add_argument('--source-docx')
    parser.add_argument('--template-dir')
    parser.add_argument('--replicate-strategy', choices=['auto', 'clone-patch', 'shell-rebuild'], default='auto')
    parser.add_argument('--audit-json')
    args = parser.parse_args()

    profile = load_json(Path(args.profile_json))
    content = load_json(Path(args.content_json))
    document_type = str(content.get('document_type') or profile.get('detected_document_type') or '').strip()
    if not document_type:
        print('ERROR: could not determine document_type from content or profile')
        return 2
    content['document_type'] = document_type

    errors, warnings = validate_word(content)
    if warnings:
        print('WARNINGS:')
        for item in warnings:
            print(f'- {item}')
    if errors:
        print('ERRORS:')
        for item in errors:
            print(f'- {item}')
        return 3

    mode = args.mode
    source_docx = Path(args.source_docx) if args.source_docx else None
    if source_docx and not source_docx.exists():
        print(f'ERROR: source DOCX not found: {source_docx}')
        return 4
    if source_docx is None:
        profile_path = Path(args.profile_json)
        template_dir = Path(args.template_dir) if args.template_dir else profile_path.parent.parent / 'assets' / 'templates'
        candidate = find_template(document_type, profile, template_dir)
        if candidate and candidate.exists():
            source_docx = candidate

    if mode == 'auto':
        mode = 'replicate' if source_docx is not None else 'canonical'

    output = Path(args.output_docx)
    audit_payload = {'mode': mode}
    if mode == 'replicate':
        if source_docx is None:
            print('ERROR: replicate mode requires --source-docx or a matching built-in template')
            return 5
        strategy = choose_replicate_strategy(profile, content, source_docx, args.replicate_strategy)
        audit_payload['replicate_strategy'] = strategy
        if strategy == 'clone-patch':
            patch_map = content_to_placeholder_map(content)
            patch_result = clone_patch_docx(source_docx, output, patch_map)
            audit_payload['patch_result'] = patch_result
            if patch_result.get('replacement_count', 0) == 0 and args.replicate_strategy == 'auto':
                generate_replicate_shell(source_docx, content, output)
                audit_payload['replicate_strategy'] = 'shell-rebuild-fallback'
                print(f'No placeholders patched from {source_docx}; fell back to shell-rebuild into {output}')
            else:
                print(f'Created {output} in replicate mode from {source_docx} using clone-patch')
        else:
            generate_replicate_shell(source_docx, content, output)
            print(f'Created {output} in replicate mode from {source_docx} using shell-rebuild')
    else:
        generate_canonical(content, output)
        print(f'Created {output} in canonical mode')

    if args.audit_json:
        Path(args.audit_json).write_text(json.dumps(audit_payload, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
