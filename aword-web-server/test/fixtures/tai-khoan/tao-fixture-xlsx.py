# Sinh tệp Excel mẫu cho test nhập danh sách tài khoản (src/tai-khoan/xlsx.ts, nhap-danh-sach.ts) bằng openpyxl.
# Chạy lại khi cần:  python test/fixtures/tai-khoan/tao-fixture-xlsx.py   — test Node chỉ đọc tệp .xlsx đã sinh.
import datetime
import os
import re
import zipfile

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904

THU_MUC = os.path.dirname(os.path.abspath(__file__))
TIEU_DE = ['Họ tên', 'Email', 'Số điện thoại', 'Mã đơn vị', 'Vai trò', 'Hạn dùng', 'Hạn mức tháng (đồng)']

# 1) Danh sách chuẩn: có dòng hợp lệ và các lỗi cần báo từng dòng
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Danh sách'
ws.append(TIEU_DE)
ws.append(['Nguyễn Văn A', 'A.Nguyen@thcs-xa.edu.vn', '0912 345 678', 'THCS01', 'Người dùng', datetime.datetime(2027, 6, 30), 200000])  # 2 hợp lệ
ws.append(['Trần Thị B', None, 912345679, 'THCS01', 'quan_tri_don_vi', '31/12/2026', '1.500.000'])  # 3 hợp lệ (ô số mất số 0 đầu)
ws.append(['Lê Thị C & "Hoa" <1>', 'le.thi.c@truong.vn', '', '', '', None, None])  # 4 hợp lệ (không đơn vị)
ws.append(['Trùng email (khác hoa thường)', 'a.nguyen@THCS-XA.edu.vn', None, 'THCS01', 'Người dùng', None, None])  # 5 trùng email dòng 2
ws.append(['Phạm Văn D', None, '+84 912 345 678', 'THCS01', 'Giáo viên', None, None])  # 6 trùng số dòng 2 + vai trò sai
ws.append(['Hoàng Thị E', 'hoang.e@truong.vn', None, 'KHONGCO', 'Người dùng', None, None])  # 7 mã đơn vị không có
ws.append(['Đã có trong hệ thống', 'dacosan@truong.vn', None, 'THCS01', 'Người dùng', None, None])  # 8 trùng CSDL
ws.append(['Số điện thoại dạng số', None, 987654321, 'THCS01', None, datetime.date(2027, 1, 15), 0])  # 9 hợp lệ
ws.append([])  # 10 dòng trống giữa bảng → bỏ qua
ws.append(['', 'sai-email', '12345', 'THCS01', 'Người dùng', 'ngày mai', '12abc'])  # 11 nhiều lỗi
ws.append(['Không có liên hệ', None, None, 'THCS01', None, None, None])  # 12 thiếu cả email và số
ws['F2'].number_format = 'dd/mm/yyyy'
wb.save(os.path.join(THU_MUC, 'nhap-tai-khoan.xlsx'))

# 2) Giả lập tệp do chính Excel lưu: openpyxl sinh nội dung (hệ ngày 1904, ô công thức), sau đó chỉnh gói ZIP cho giống
#    Excel — chuỗi dùng chung (sharedStrings.xml, có ô chữ nhiều đoạn + phiên âm <rPh>), công thức có giá trị lưu sẵn,
#    đường dẫn quan hệ tương đối, trang đầu tiên nằm ở sheet2.xml, thẻ có tiền tố không gian tên (x:).
wb = openpyxl.Workbook()
wb.epoch = CALENDAR_MAC_1904
chinh = wb.active
chinh.title = 'Nhập'
chinh.append(TIEU_DE)
chinh.append(['Đinh Văn G', 'dinh.g@truong.vn', None, 'THCS01', 'Người dùng', datetime.datetime(2026, 12, 31), 350000])
chinh['A3'] = '=CONCATENATE("Mai ","Thị H")'
chinh['C3'] = '0977000111'
chinh['A4'] = 'GIU_CHO_CHU_NHIEU_DOAN'
chinh['B4'] = 'bui.k@truong.vn'
phu = wb.create_sheet('Ghi chú')
phu['A1'] = 'Trang này không được đọc'
tam = os.path.join(THU_MUC, '_tam.xlsx')
wb.save(tam)

with zipfile.ZipFile(tam) as z:
    tep = {i.filename: z.read(i.filename) for i in z.infolist()}
os.remove(tam)

chuoi = []


def sang_chuoi_chung(xml):
    def thay(m):
        van = m.group(2)
        if van == 'GIU_CHO_CHU_NHIEU_DOAN':
            chuoi.append('<si><r><t xml:space="preserve">Bùi </t></r><r><rPr><b/></rPr><t>Thị K</t></r>'
                         '<rPh sb="0" eb="1"><t>PHIEN_AM_KHONG_LAY</t></rPh></si>')
        else:
            chuoi.append('<si><t>' + van + '</t></si>')
        return '<c r="%s" t="s"><v>%d</v></c>' % (m.group(1), len(chuoi) - 1)
    return re.sub(r'<c r="([A-Z]+\d+)" t="inlineStr"><is><t>(.*?)</t></is></c>', thay, xml)


trang_nhap = sang_chuoi_chung(tep['xl/worksheets/sheet1.xml'].decode('utf-8'))
trang_nhap = trang_nhap.replace('<c r="A3"><f>', '<c r="A3" t="str"><f>').replace('<v></v></c>', '<v>Mai Th&#7883; H</v></c>')
trang_nhap = re.sub(r'<(/?)([A-Za-z]+)([ >/])', r'<\1x:\2\3', trang_nhap)
trang_nhap = trang_nhap.replace('<x:worksheet xmlns=', '<x:worksheet xmlns:x=')
trang_ghi_chu = sang_chuoi_chung(tep['xl/worksheets/sheet2.xml'].decode('utf-8'))
tep['xl/worksheets/sheet2.xml'] = trang_nhap.encode('utf-8')
tep['xl/worksheets/sheet1.xml'] = trang_ghi_chu.encode('utf-8')
tep['xl/sharedStrings.xml'] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
                               '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d" uniqueCount="%d">%s</sst>'
                               % (len(chuoi), len(chuoi), ''.join(chuoi))).encode('utf-8')
rels = tep['xl/_rels/workbook.xml.rels'].decode('utf-8')
rels = rels.replace('Target="/xl/worksheets/sheet1.xml"', 'Target="worksheets/TAM"')
rels = rels.replace('Target="/xl/worksheets/sheet2.xml"', 'Target="worksheets/sheet1.xml"').replace('worksheets/TAM', 'worksheets/sheet2.xml')
rels = rels.replace('</Relationships>', '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
tep['xl/_rels/workbook.xml.rels'] = rels.encode('utf-8')
tep['[Content_Types].xml'] = tep['[Content_Types].xml'].replace(
    b'</Types>', b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
with zipfile.ZipFile(os.path.join(THU_MUC, 'nhap-kieu-excel-1904.xlsx'), 'w', zipfile.ZIP_DEFLATED) as z:
    for ten in ['[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', 'xl/styles.xml',
                'xl/sharedStrings.xml', 'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml', 'xl/theme/theme1.xml',
                'docProps/core.xml', 'docProps/app.xml']:
        z.writestr(ten, tep[ten])
print('Đã ghi các tệp .xlsx vào', THU_MUC)
